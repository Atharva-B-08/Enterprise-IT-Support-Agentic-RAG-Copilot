from pathlib import Path
from typing import Iterable, List
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter 
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from docx import Document as DocxDocument   # this is used for loading .docx files

SUPPORTED_FILE_TYPES = {"*.pdf", "*.txt", "*.md", "*.docx"}

def load_file(file_path: Path) -> list[Document]:
    """
    Load a file and return a list of Document objects.
    Supports PDF, TXT, MD, and DOCX files.
    """
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        return PyPDFLoader(str(file_path)).load()
    if suffix in {".txt", ".md"}:
        return TextLoader(str(file_path), encoding="utf-8").load()
    if suffix == ".docx":
        docs = DocxDocument(str(file_path))
        text = "\n".join(p.text for p in docs.paragraphs if p.text.strip())
        return [Document(page_content=text, metadata={"source": str(file_path)})]
    raise ValueError(f"Unsupported file type: {suffix}")


def chunk_documents(docs: Iterable[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=120, add_start_index=True)
    return splitter.split_documents(list(docs))
    